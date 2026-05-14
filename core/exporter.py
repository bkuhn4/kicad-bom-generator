from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

EXPORT_COLUMNS = [
    ("qty",          "Quantity"),
    ("references",   "References"),
    ("value",        "Value"),
    ("description",  "Description"),
    ("footprint",    "Footprint"),
    ("packaging",    "Packaging"),
    ("mpn",          "MPN"),
    ("manufacturer", "Manufacturer"),
    ("digikey_pn",   "DigiKey PN"),
    ("stock",        "DigiKey Stock"),
    ("mouser_pn",    "Mouser PN"),
    ("lcsc_pn",      "LCSC PN"),
    ("notes",        "Notes"),
]

_REFS_MAX_WIDTH = 24


# ── color helpers ────────────────────────────────────────────────────────────

def _lighten(hex6: str, factor: float) -> str:
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    r = min(255, int(r + (255 - r) * factor))
    g = min(255, int(g + (255 - g) * factor))
    b = min(255, int(b + (255 - b) * factor))
    return f"{r:02X}{g:02X}{b:02X}"


def _text_on(hex6: str) -> str:
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    return "FFFFFF" if (0.299 * r + 0.587 * g + 0.114 * b) < 140 else "000000"


def _hide_table_filter_buttons(filepath: str | Path):
    """Post-process xlsx to set showAutoFilter=0 on every Table element."""
    buf = io.BytesIO()
    with zipfile.ZipFile(filepath, "r") as zin:
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.startswith("xl/tables/") and info.filename.endswith(".xml"):
                    data = re.sub(rb' showAutoFilter="[^"]*"', b"", data)
                    data = data.replace(b"<table ", b'<table showAutoFilter="0" ', 1)
                zout.writestr(info, data)
    Path(filepath).write_bytes(buf.getvalue())


# ── main export ──────────────────────────────────────────────────────────────

def export_bom(rows: list[dict], filepath: str | Path, header_color: str = "70AD47",
               columns: list[tuple] | None = None,
               project_name: str = "", revision: str = "",
               show_title: bool = True, show_footer: bool = True):
    cols = columns if columns is not None else EXPORT_COLUMNS
    n_cols = len(cols)
    hex6 = header_color.lstrip("#").upper()
    stripe_hex = _lighten(hex6, 0.72)

    bside = Side(style="thin", color=hex6)
    cell_border = Border(left=bside, right=bside, top=bside, bottom=bside)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    # ── narrow page margins ──────────────────────────────────────
    ws.page_margins.left   = 0.25
    ws.page_margins.right  = 0.25
    ws.page_margins.top    = 0.75
    ws.page_margins.bottom = 0.75
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3

    # ── page header / footer ─────────────────────────────────────
    if show_title and (project_name or revision):
        if project_name:
            ws.oddHeader.center.text = f"{project_name} BOM"
            ws.oddHeader.center.font = "Bold"
            ws.oddHeader.center.size = 12
        if revision:
            ws.oddHeader.right.text = f"Rev {revision}"
            ws.oddHeader.right.font = "Bold"
            ws.oddHeader.right.size = 11

    if show_footer:
        now = datetime.now()
        t = now.strftime("%I:%M %p").lstrip("0")
        date_str = f"{now.month}/{now.day}/{now.year}"
        ws.oddFooter.left.text = f"Created: {t} {date_str}"
        ws.oddFooter.left.size = 9

    # ── column header row ────────────────────────────────────────
    hdr_fill  = PatternFill("solid", fgColor=hex6)
    hdr_font  = Font(bold=True, color=_text_on(hex6), size=11)
    hdr_align = Alignment(horizontal="left", vertical="center")

    for col, (_, label) in enumerate(cols, 1):
        c = ws.cell(row=1, column=col, value=label)
        c.fill      = hdr_fill
        c.font      = hdr_font
        c.alignment = hdr_align
        c.border    = cell_border

    ws.row_dimensions[1].height = 22

    # ── data rows ────────────────────────────────────────────────
    stripe_fill  = PatternFill("solid", fgColor=stripe_hex)
    white_fill   = PatternFill("solid", fgColor="FFFFFF")
    data_align   = Alignment(vertical="center", wrap_text=False)
    wrap_align   = Alignment(vertical="top",    wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    data_font    = Font(size=10)

    for ri, row in enumerate(rows, 2):
        fill = stripe_fill if ri % 2 == 0 else white_fill
        for ci, (key, _) in enumerate(cols, 1):
            c = ws.cell(row=ri, column=ci, value=row.get(key, ""))
            c.fill   = fill
            c.border = cell_border
            c.font   = data_font
            if key == "qty":
                c.alignment = center_align
            elif key == "references":
                c.alignment = wrap_align
            else:
                c.alignment = data_align

    # ── column widths ────────────────────────────────────────────
    for ci, (key, _) in enumerate(cols, 1):
        col_letter = get_column_letter(ci)
        max_len = max(
            (len(str(ws.cell(row=r, column=ci).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        auto_w = min(max_len * 1.1 + 4, 80)
        if key == "references":
            ws.column_dimensions[col_letter].width = min(auto_w, _REFS_MAX_WIDTH)
        else:
            ws.column_dimensions[col_letter].width = auto_w

    # ── Excel Table ──────────────────────────────────────────────
    last_col = get_column_letter(n_cols)
    tab = Table(displayName="BOM", ref=f"A1:{last_col}{1 + len(rows)}")
    ws.add_table(tab)

    ws.freeze_panes = "A2"
    wb.save(filepath)

    _hide_table_filter_buttons(filepath)


def export_csv(rows: list[dict], filepath: str | Path, columns: list[tuple] | None = None):
    cols = columns if columns is not None else EXPORT_COLUMNS
    keys = [k for k, _ in cols]
    labels = [l for _, l in cols]

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(labels)
        for row in rows:
            writer.writerow([row.get(k, "") for k in keys])

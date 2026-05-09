from __future__ import annotations
from typing import List

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import BomRow
from . import config

_HEADER_FILL = PatternFill("solid", fgColor="70AD47")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_EVEN_FILL   = PatternFill("solid", fgColor="E2EFDA")
_ODD_FILL    = PatternFill("solid", fgColor="FFFFFF")

_BORDER = Border(
    left=Side(border_style="thin", color="70AD47"),
    right=Side(border_style="thin", color="70AD47"),
    top=Side(border_style="thin", color="70AD47"),
    bottom=Side(border_style="thin", color="70AD47")
)

HEADERS = [
    "Reference(s)", "Description", "Package", "MF Part Number",
    "Manufacturer", "DigiKey PN", "DK Packaging",
    "Mouser PN", "Mouser Packaging", "LCSC PN", "Notes",
]

_REF_COL_WIDTH = 28   # References column: fixed
_MIN_WIDTH     = 8
_MAX_WIDTH     = 55


def export(rows: List[BomRow], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    
    cfg = config.load()
    export_columns = cfg.get("export_columns", {h: True for h in HEADERS})

    # Filter headers based on config
    active_headers = [h for h in HEADERS if export_columns.get(h, True)]
    
    ws.append(active_headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = _BORDER

    # Track max content length per column to auto-size
    col_widths = [len(h) + 2 for h in active_headers]

    for i, row in enumerate(rows, start=2):
        row_data_dict = {
            "Reference(s)": row.references,
            "Description": row.description,
            "Package": row.package,
            "MF Part Number": row.mf_part_number,
            "Manufacturer": row.manufacturer,
            "DigiKey PN": row.digikey_pn,
            "DK Packaging": row.digikey_selected_packaging,
            "Mouser PN": row.mouser_pn,
            "Mouser Packaging": row.mouser_selected_packaging,
            "LCSC PN": row.lcsc_pn,
            "Notes": row.notes,
        }
        
        data = [row_data_dict[h] for h in active_headers]
        ws.append(data)

        fill = _EVEN_FILL if i % 2 == 0 else _ODD_FILL
        for j, cell in enumerate(ws[i]):
            cell.fill = fill
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=(active_headers[j] == "Reference(s)"), 
            )
            cell.border = _BORDER
            cell_len = len(str(data[j])) if data[j] else 0
            col_widths[j] = max(col_widths[j], cell_len + 2)

    for j, width in enumerate(col_widths):
        letter = get_column_letter(j + 1)
        if active_headers[j] == "Reference(s)":
            ws.column_dimensions[letter].width = _REF_COL_WIDTH
        else:
            ws.column_dimensions[letter].width = max(_MIN_WIDTH, min(_MAX_WIDTH, width))

    ws.freeze_panes = "A2"
    wb.save(path)

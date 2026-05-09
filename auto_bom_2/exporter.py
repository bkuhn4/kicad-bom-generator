from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORT_COLUMNS = [
    ("qty",          "Quantity"),
    ("description",  "Description"),
    ("mpn",          "Manufacturer PN"),
    ("manufacturer", "Manufacturer"),
    ("digikey_pn",   "Digikey PN"),
    ("mouser_pn",    "Mouser PN"),
    ("lcsc_pn",      "LCSC PN"),
    ("notes",        "Notes"),
]


# ── color helpers ────────────────────────────────────────────────────────────

def _lighten(hex6: str, factor: float) -> str:
    """Blend hex6 toward white by factor (0 = unchanged, 1 = white)."""
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    r = min(255, int(r + (255 - r) * factor))
    g = min(255, int(g + (255 - g) * factor))
    b = min(255, int(b + (255 - b) * factor))
    return f"{r:02X}{g:02X}{b:02X}"


def _text_on(hex6: str) -> str:
    """Return FFFFFF or 000000 for legible text on a background."""
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    return "FFFFFF" if (0.299 * r + 0.587 * g + 0.114 * b) < 140 else "000000"


# ── main export ──────────────────────────────────────────────────────────────

def export_bom(rows: list[dict], filepath: str | Path, header_color: str = "4F6228"):
    hex6 = header_color.lstrip("#").upper()
    stripe_hex  = _lighten(hex6, 0.72)   # light tint for alternating rows
    text_color  = _text_on(hex6)

    # All borders use the header color so the grid lines match
    bside = Side(style="thin", color=hex6)
    cell_border = Border(left=bside, right=bside, top=bside, bottom=bside)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    # ── header row ───────────────────────────────────────────────
    hdr_fill  = PatternFill("solid", fgColor=hex6)
    hdr_font  = Font(bold=True, color=text_color, size=11)
    hdr_align = Alignment(horizontal="center", vertical="center")

    for col, (_, label) in enumerate(EXPORT_COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=label)
        c.fill      = hdr_fill
        c.font      = hdr_font
        c.alignment = hdr_align
        c.border    = cell_border

    ws.row_dimensions[1].height = 22

    # ── data rows ────────────────────────────────────────────────
    stripe_fill = PatternFill("solid", fgColor=stripe_hex)
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    data_align  = Alignment(vertical="center", wrap_text=False)
    data_font   = Font(size=10)

    for ri, row in enumerate(rows, 2):
        fill = stripe_fill if ri % 2 == 0 else white_fill
        for ci, (key, _) in enumerate(EXPORT_COLUMNS, 1):
            c = ws.cell(row=ri, column=ci, value=row.get(key, ""))
            c.fill      = fill
            c.border    = cell_border
            c.alignment = data_align
            c.font      = data_font

    # ── auto-filter (filter-arrow dropdowns) ────────────────────
    last_col = get_column_letter(len(EXPORT_COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    # ── auto-width ───────────────────────────────────────────────
    for ci in range(1, len(EXPORT_COLUMNS) + 1):
        col_letter = get_column_letter(ci)
        max_len = max(
            (len(str(ws.cell(row=r, column=ci).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"
    wb.save(filepath)
